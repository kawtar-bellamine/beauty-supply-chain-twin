"""Tests on the committed sample of the real data (about 30 seconds, no download)."""
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from bsc import excel_model, flow, forecast, twin
from bsc.data import build
from bsc.leadtime import FreightModel, TransitModel


@pytest.fixture(scope="module")
def con():
    return build(sample=True)


@pytest.fixture(scope="module")
def frames(con):
    return (con.execute("SELECT * FROM fct_orders").df(), con.execute("SELECT km, freight, is_beauty FROM "
                                                                      "fct_order_lines").df(),
            con.execute("SELECT * FROM daily_beauty_demand").df())


@pytest.fixture(scope="module")
def setup(frames):
    orders, lines, _ = frames
    tm, fm = TransitModel.fit(orders), FreightModel.fit(lines)
    real = twin.prepare_orders(orders)
    return tm, fm, twin.promise_days(real, tm), twin.brand_scale(real, volume=150, seed=0)


def test_one_row_per_order_and_coordinates(con):
    n, u = con.execute("SELECT count(*), count(DISTINCT order_id) FROM fct_orders").fetchone()
    assert n == u
    assert con.execute("SELECT avg((km IS NULL)::INT) FROM fct_orders").fetchone()[0] < 0.03


def test_stage_times_add_up(frames):
    o = frames[0].dropna(subset=["d_approval", "d_handling", "d_transit"])
    o = o[o["d_handling"] >= 0]
    assert np.allclose(o["d_approval"] + o["d_handling"] + o["d_transit"], o["d_total"], atol=1e-6)


def test_late_causes_are_a_partition(frames):
    c = flow.late_causes(flow.beauty(frames[0]))
    assert c["late_caused_by_carrier_only_pct"] + c["late_with_seller_delay_pct"] == pytest.approx(100)
    assert c["late_if_shipped_late_pct"] > c["late_if_shipped_on_time_pct"]


def test_seller_pareto(frames):
    p = flow.seller_pareto(flow.beauty(frames[0]))
    assert p["cum_share_pct"].is_monotonic_increasing and p["cum_share_pct"].iloc[-1] == pytest.approx(100)


def test_transit_grows_with_distance(setup):
    tm = setup[0]
    assert pd.Series(tm.p50).rank().corr(pd.Series(range(len(tm.p50)), dtype=float)) > 0.8
    s = tm.sample(np.array([50.0, 2500.0] * 50), np.random.default_rng(0))
    assert s[1::2].mean() > s[::2].mean()


def test_forecast_uses_only_the_past(frames):
    w = forecast.weekly(frames[2])
    a = forecast.backtest(w, horizon=2, n_origins=6)
    w2 = w.copy()
    w2.iloc[-1] *= 10
    b = forecast.backtest(w2, horizon=2, n_origins=6)
    first = a["origin"].min()
    pd.testing.assert_series_equal(a[a.origin == first]["forecast"].reset_index(drop=True),
                                   b[b.origin == first]["forecast"].reset_index(drop=True))


def test_twin_serves_every_order(setup):
    tm, fm, prom, stream = setup
    r = twin.run(stream, tm, fm, prom)
    o = r["_orders"]
    assert len(o) == len(stream) and (o["wait"] >= 0).all()
    assert 0 <= r["fill_rate_pct"] <= 100 and r["cost_per_unit"] > 0


def test_more_capacity_never_hurts(setup):
    tm, fm, prom, stream = setup
    lo = twin.run(stream, tm, fm, prom, capacity_ratio=1.1)
    hi = twin.run(stream, tm, fm, prom, capacity_ratio=2.0)
    assert hi["fill_rate_pct"] >= lo["fill_rate_pct"]


def test_regional_dc_shortens_north_east_deliveries(setup):
    tm, fm, prom, stream = setup
    one = twin.run(stream, tm, fm, prom)["_orders"]
    two = twin.run(stream, tm, fm, prom, dcs=("São Paulo DC", "Recife DC"))["_orders"]
    assert two[two.region == "NE"]["lead"].median() < one[one.region == "NE"]["lead"].median()


def test_excel_model_has_live_formulas(tmp_path):
    p = excel_model.build(tmp_path / "m.xlsx", [{"name": "A", "mean": 100, "std": 20, "km": 0},
                                                {"name": "B", "mean": 10, "std": 5, "km": 2000}],
                          [{"scenario": "x", "fill_rate_pct": 95.0}],
                          {"service": 0.95, "unit_value": 80, "holding_rate": 0.25, "kmpd": 650, "dc_fixed": 1500})
    ws = load_workbook(p)["Inventory policy"]
    assert str(ws["F4"].value).startswith("=") and "SQRT" in ws["F4"].value
    assert ws["E4"].value == "=NORMSINV(Inputs!$B$4)"
