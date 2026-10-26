"""Excel inventory-policy model with live formulas: change a blue input, every result recalculates."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BLUE = Font(name="Arial", color="0000FF", size=10)
BLACK = Font(name="Arial", color="000000", size=10)
GREEN = Font(name="Arial", color="008000", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
TITLE = Font(name="Arial", bold=True, size=14)
HEAD = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="1A1A1A")
GOLD = PatternFill("solid", fgColor="F3E7D3")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))


def _head(ws, row, labels, col=1):
    for i, h in enumerate(labels):
        c = ws.cell(row, col + i, h)
        c.font, c.fill = HEAD, HFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build(path: Path, dcs: list[dict], twin_rows: list[dict], params: dict) -> Path:
    """dcs: [{name, mean, std, km}]; twin_rows: scenario results from the simulation (values)."""
    wb = Workbook()
    rd = wb.active
    rd.title = "Read me"
    rd["A1"] = "Beauty supply chain: inventory policy model"
    rd["A1"].font = TITLE
    notes = [
        "How to use: change the blue cells on the Inputs sheet (yellow = key levers); every other number is a formula.",
        "Colours: blue = input, black = formula, green = link to another sheet.",
        "Demand per DC: mean and standard deviation of daily units, January to August 2018, from the simulation of "
        "the real Olist beauty order stream scaled to a brand-sized flow (see the project README).",
        "Line-haul lead time = 1 day of loading + distance / truck km per day, rounded up.",
        "Safety stock = z x sigma x SQRT(lead time + review period); z = NORMSINV(cycle service level).",
        "Average stock = safety stock + half a review period of demand (cycle stock) + pipeline is excluded.",
        "Twin results sheet: values copied from the Python simulation (experiments/run_all.py), not formulas.",
    ]
    for i, n in enumerate(notes, 3):
        rd.cell(i, 1, n).font = BLACK
    rd.column_dimensions["A"].width = 120

    # ------------------------------------------------------------------ inputs
    ws = wb.create_sheet("Inputs")
    ws["A1"] = "Inputs"
    ws["A1"].font = TITLE
    glob = [("Cycle service level", params["service"], "0.0%", "Target probability of no stockout per cycle"),
            ("Review period (days)", 1, "0", "DCs review stock every day"),
            ("Demand growth vs 2018 level", 0.0, "0.0%", "Scenario lever: +30% tests a surge"),
            ("Unit value (BRL)", params["unit_value"], "#,##0.00", "Median beauty item price in the data"),
            ("Holding rate per year", params["holding_rate"], "0.0%", "Assumption: cost of capital, storage, obsolescence"),
            ("Truck km per day", params["kmpd"], "#,##0", "Assumption: full-truck line haul"),
            ("Regional DC fixed cost per day (BRL)", params["dc_fixed"], "#,##0", "Assumption: rent and staff")]
    _head(ws, 3, ["Parameter", "Value", "Note"])
    for r, (k, v, fmt, note) in enumerate(glob, 4):
        ws.cell(r, 1, k).font = BLACK
        c = ws.cell(r, 2, v)
        c.font, c.number_format = BLUE, fmt
        if r in (4, 6):
            c.fill = YELLOW
        ws.cell(r, 3, note).font = BLACK
    dv = DataValidation(type="decimal", operator="between", formula1="0.5", formula2="0.999")
    ws.add_data_validation(dv)
    dv.add("B4")
    _head(ws, 13, ["Distribution centre", "Mean daily demand (units)", "Std daily demand (units)",
                   "Line-haul distance from plant (km)"])
    for r, d in enumerate(dcs, 14):
        ws.cell(r, 1, d["name"]).font = BLACK
        for j, key in enumerate(["mean", "std", "km"], 2):
            c = ws.cell(r, j, round(float(d[key]), 1))
            c.font, c.number_format = BLUE, "#,##0.0"
    ws.cell(14, 6).comment = Comment("Source: simulation of the real order stream, 2018-01 to 2018-08.", "model")
    for col, w in zip("ABCD", [38, 22, 22, 26]):
        ws.column_dimensions[col].width = w

    # ------------------------------------------------------------------ policy
    po = wb.create_sheet("Inventory policy")
    po["A1"] = "Order-up-to policy per DC"
    po["A1"].font = TITLE
    cols = ["Distribution centre", "Mean demand incl. growth", "Lead time (days)", "Protection period (days)", "z",
            "Safety stock (units)", "Order-up-to level S (units)", "Average stock (units)", "Days of cover",
            "Holding cost per year (BRL)"]
    _head(po, 3, cols)
    n = len(dcs)
    for i in range(n):
        r, s = 4 + i, 14 + i
        po.cell(r, 1, f"=Inputs!A{s}").font = GREEN
        po.cell(r, 2, f"=Inputs!B{s}*(1+Inputs!$B$6)").font = GREEN
        po.cell(r, 3, f"=1+ROUNDUP(Inputs!D{s}/Inputs!$B$9,0)").font = GREEN
        po.cell(r, 4, f"=C{r}+Inputs!$B$5")
        po.cell(r, 5, "=NORMSINV(Inputs!$B$4)")
        po.cell(r, 6, f"=E{r}*Inputs!C{s}*(1+Inputs!$B$6)*SQRT(D{r})")
        po.cell(r, 7, f"=B{r}*D{r}+F{r}")
        po.cell(r, 8, f"=F{r}+B{r}*Inputs!$B$5/2")
        po.cell(r, 9, f"=IF(B{r}>0,H{r}/B{r},0)")
        po.cell(r, 10, f"=H{r}*Inputs!$B$7*Inputs!$B$8")
        for j in range(2, 11):
            cell = po.cell(r, j)
            cell.number_format = "0.00" if j in (5, 9) else "#,##0"
            if cell.font != GREEN:
                cell.font = BLACK if j > 3 else cell.font
    tr = 4 + n
    po.cell(tr, 1, "Total").font = BOLD
    for j, col in [(6, "F"), (8, "H"), (10, "J")]:
        c = po.cell(tr, j, f"=SUM({col}4:{col}{tr - 1})")
        c.font, c.number_format, c.fill = BOLD, "#,##0", GOLD
    for col, w in zip("ABCDEFGHIJ", [24, 16, 12, 14, 8, 14, 16, 14, 10, 18]):
        po.column_dimensions[col].width = w

    # ------------------------------------------------------------------ service level trade-off
    sv = wb.create_sheet("Service trade-off")
    sv["A1"] = "Safety stock and holding cost versus cycle service level (all DCs)"
    sv["A1"].font = TITLE
    _head(sv, 3, ["Cycle service level", "z", "Total safety stock (units)", "Holding cost of safety stock (BRL/yr)"])
    levels = [0.80, 0.85, 0.90, 0.95, 0.975, 0.99, 0.995, 0.999]
    for i, lv in enumerate(levels, 4):
        c = sv.cell(i, 1, lv)
        c.font, c.number_format = BLUE, "0.0%"
        sv.cell(i, 2, f"=NORMSINV(A{i})").number_format = "0.00"
        terms = "+".join(f"Inputs!C{14 + k}*SQRT('Inventory policy'!D{4 + k})" for k in range(n))
        sv.cell(i, 3, f"=B{i}*({terms})*(1+Inputs!$B$6)").number_format = "#,##0"
        sv.cell(i, 4, f"=C{i}*Inputs!$B$7*Inputs!$B$8").number_format = "#,##0"
    ch = LineChart()
    ch.title = "Holding cost of safety stock vs service level"
    ch.y_axis.title = "BRL per year"
    ch.x_axis.title = "Cycle service level"
    ch.add_data(Reference(sv, min_col=4, min_row=3, max_row=3 + len(levels)), titles_from_data=True)
    ch.set_categories(Reference(sv, min_col=1, min_row=4, max_row=3 + len(levels)))
    ch.height, ch.width = 8, 16
    sv.add_chart(ch, "F3")
    for col, w in zip("ABCD", [20, 10, 24, 30]):
        sv.column_dimensions[col].width = w

    # ------------------------------------------------------------------ twin results (values)
    tw = wb.create_sheet("Twin results")
    tw["A1"] = "Scenario results from the Python digital twin (values, not formulas)"
    tw["A1"].font = TITLE
    keys = list(twin_rows[0].keys())
    _head(tw, 3, keys)
    for i, row in enumerate(twin_rows, 4):
        for j, k in enumerate(keys, 1):
            v = row[k]
            c = tw.cell(i, j, round(v, 2) if isinstance(v, float) else v)
            c.font = BLACK
            c.border = THIN
    for j in range(1, len(keys) + 1):
        tw.column_dimensions[tw.cell(3, j).column_letter].width = 18 if j > 1 else 40
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.font is not None and c.font.name != "Arial":
                    c.font = Font(name="Arial", size=c.font.size or 10, bold=c.font.bold, color=c.font.color)
    wb.save(path)
    return path
